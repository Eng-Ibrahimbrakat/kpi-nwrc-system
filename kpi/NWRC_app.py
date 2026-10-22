import pandas as pd
import os
import gspread
from oauth2client.service_account import ServiceAccountCredentials
import plotly.express as px
from io import BytesIO
import streamlit as st
from datetime import datetime, timedelta
import time

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============================================================
# إعداد الصفحة
# ============================================================
st.set_page_config(page_title="نظام مؤشرات الأداء - المركز القومي لبحوث المياه",
                   layout="wide", page_icon="📊")

current_dir = os.path.dirname(__file__)
logo_path = os.path.join(current_dir, "logo.png")
if os.path.exists(logo_path):
    st.image(logo_path, width=150)

st.markdown("""
<style>
html, body, [class*="css"] { direction: rtl; text-align: right; }
.stButton > button { width: 100%; }
[data-testid="stMetricValue"] { font-size: 1.5rem; }
[data-testid="stDataFrame"], [data-testid="stDataEditor"] { direction: rtl !important; }
[data-testid="stDataFrame"] table, [data-testid="stDataEditor"] table { direction: rtl !important; text-align: right !important; }
[data-testid="stDataFrame"] th, [data-testid="stDataFrame"] td,
[data-testid="stDataEditor"] th, [data-testid="stDataEditor"] td { text-align: right !important; direction: rtl !important; }
h1, h2, h3, h4, h5, h6 { direction: rtl; text-align: right; }
.stTabs [data-baseweb="tab-list"] { direction: rtl; }
.stTextInput input, .stNumberInput input, .stTextArea textarea { direction: rtl; text-align: right; }
.stSelectbox div[data-baseweb="select"] { direction: rtl; }
</style>
""", unsafe_allow_html=True)

def get_local_time():
    return (datetime.now() + timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S")

# ============================================================
# Google Sheets
# ============================================================
SPREADSHEET_KEY = "1QSfmNo9U0TNvdwRgLhLBgVNZbiL8wVcoWlffBz6cSfg"

@st.cache_resource(show_spinner=False)
def _get_client():
    scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
    creds = ServiceAccountCredentials.from_json_keyfile_dict(st.secrets["gcp_service_account"], scope)
    return gspread.authorize(creds)

@st.cache_resource(show_spinner=False)
def _get_spreadsheet_cached(_bust=0):
    return _get_client().open_by_key(SPREADSHEET_KEY)

def get_spreadsheet():
    return _get_spreadsheet_cached(_bust=st.session_state.get("gs_bust", 0))

@st.cache_resource(show_spinner=False)
def _get_ws_cached(sheet_name, _bust=0):
    try:
        return _get_spreadsheet_cached(_bust=_bust).worksheet(sheet_name)
    except Exception:
        return None

def get_worksheet(sheet_name):
    return _get_ws_cached(sheet_name, _bust=st.session_state.get("gs_bust", 0))

def get_or_create_sheet(sheet_name, headers=None):
    ws = get_worksheet(sheet_name)
    if ws is not None:
        return ws
    time.sleep(0.5)
    try:
        return get_spreadsheet().worksheet(sheet_name)
    except Exception:
        pass
    try:
        ws = get_spreadsheet().add_worksheet(title=sheet_name, rows=2000, cols=40)
        if headers:
            ws.append_row(headers)
        st.session_state.gs_bust = st.session_state.get("gs_bust", 0) + 1
        return ws
    except Exception as e:
        if "already exists" in str(e).lower():
            time.sleep(1)
            try:
                return get_spreadsheet().worksheet(sheet_name)
            except Exception:
                return None
        st.error(f"خطأ في الوصول للورقة {sheet_name}: {e}")
        return None

@st.cache_data(ttl=120, show_spinner=False)
def _fetch_cached(sheet_name, _bust=0):
    try:
        ws = get_worksheet(sheet_name)
        if ws is None:
            return []
        for attempt in range(3):
            try:
                return ws.get_all_values()
            except gspread.exceptions.APIError as e:
                if "429" in str(e) or "Quota" in str(e):
                    time.sleep(2 ** attempt)
                    continue
                raise
        return []
    except Exception:
        return []

def fetch_data(sheet_name):
    return _fetch_cached(sheet_name, _bust=st.session_state.get("data_bust", 0))

def invalidate_all():
    st.session_state.gs_bust = st.session_state.get("gs_bust", 0) + 1
    st.session_state.data_bust = st.session_state.get("data_bust", 0) + 1
    _fetch_cached.clear()
    for key in list(st.session_state.keys()):
        if key.startswith("_cd_") or key.startswith("_cb_"):
            del st.session_state[key]

def reset_all():
    invalidate_all()
    for key in list(st.session_state.keys()):
        if key.startswith("_saved_"):
            del st.session_state[key]

def fix_encoding(t):
    if isinstance(t, str):
        try:
            return t.encode('latin1').decode('utf-8')
        except Exception:
            return t
    return t

def rtl_display(df):
    if df.empty:
        return df
    return df[list(reversed(df.columns))]

def load_sheet(sheet_name, columns):
    cache_key = f"_cd_{sheet_name}"
    bust_key = f"_cb_{sheet_name}"
    cur_bust = st.session_state.get("data_bust", 0)
    if st.session_state.get(bust_key, -1) == cur_bust and cache_key in st.session_state:
        return st.session_state[cache_key].copy()
    get_or_create_sheet(sheet_name, columns)
    all_values = fetch_data(sheet_name)
    if len(all_values) <= 1:
        df = pd.DataFrame(columns=columns)
    else:
        headers = [fix_encoding(h) for h in all_values[0]]
        max_len = len(headers)
        normalized = []
        for row in all_values[1:]:
            row = list(row) + [""] * (max_len - len(row))
            normalized.append(row[:max_len])
        df = pd.DataFrame(normalized, columns=headers)
        for col in df.columns:
            if df[col].dtype == 'object' and col not in columns:
                df[col] = df[col].apply(fix_encoding)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    st.session_state[cache_key] = df.copy()
    st.session_state[bust_key] = cur_bust
    return df

def save_list(sheet_name, input_df, base_cols, institute, period_key, period_val):
    try:
        full_cols = base_cols + ["المعهد", period_key]
        ws = get_or_create_sheet(sheet_name, full_cols)
        if ws is None:
            st.error("فشل الاتصال")
            return False
        all_values = ws.get_all_values()
        kept = []
        if len(all_values) > 1:
            existing = pd.DataFrame(all_values[1:], columns=all_values[0])
            if "المعهد" in existing.columns and period_key in existing.columns:
                mask = ~((existing["المعهد"] == institute) & (existing[period_key] == period_val))
                kept = existing[mask].values.tolist()
        ws.clear()
        ws.append_row(full_cols)
        if kept:
            ws.append_rows([[str(v) for v in r] for r in kept])
        new_rows = []
        for _, row in input_df.iterrows():
            vals = [str(v) if pd.notna(v) else "" for v in row.tolist()]
            if all(v.strip() in ["", "nan", "None"] for v in vals):
                continue
            new_rows.append(vals + [institute, period_val])
        if new_rows:
            ws.append_rows(new_rows)
        clean_df = input_df.copy().reset_index(drop=True)
        save_key = f"_saved_{sheet_name}_{institute}_{period_val}"
        st.session_state[save_key] = clean_df
        st.session_state["last_save_time"] = get_local_time()
        st.session_state.data_bust = st.session_state.get("data_bust", 0) + 1
        _fetch_cached.clear()
        for key in list(st.session_state.keys()):
            if key.startswith("_cd_") or key.startswith("_cb_"):
                del st.session_state[key]
        return True
    except Exception as e:
        st.error(f"خطأ في الحفظ: {e}")
        return False

def save_or_update_row(sheet_name, all_cols, data_dict, key_cols):
    try:
        ws = get_or_create_sheet(sheet_name, all_cols)
        if ws is None:
            return None
        all_v = ws.get_all_values()
        if len(all_v) == 0:
            ws.append_rows([list(all_cols), [str(data_dict.get(c, "0")) for c in all_cols]], value_input_option="USER_ENTERED")
            return "added"
        headers = all_v[0]
        for c in all_cols:
            if c not in headers:
                headers.append(c)
        new_data = [headers]
        found = False
        for row in all_v[1:]:
            row_list = (list(row) + [""] * len(headers))[:len(headers)]
            is_match = all(k in headers and str(row_list[headers.index(k)]) == str(data_dict.get(k, "")) for k in key_cols)
            if is_match:
                new_data.append([str(data_dict.get(h, "0")) for h in headers])
                found = True
            else:
                new_data.append(row_list)
        if not found:
            new_data.append([str(data_dict.get(h, "0")) for h in headers])
        ws.clear()
        ws.append_rows(new_data, value_input_option="USER_ENTERED")
        return "updated" if found else "added"
    except Exception as e:
        st.error(f"خطأ في الحفظ: {e}")
        return None

# ============================================================
# دوال RTL لـ Word
# ============================================================
def _rtl_run(run, bold=False, size=11):
    rPr = run._element.get_or_add_rPr()
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)
    run.font.name = 'Arial'; run.font.size = Pt(size); run.font.bold = bold

def _rtl_para(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def _rtl_table(table):
    table._tbl.tblPr.append(OxmlElement('w:bidiVisual'))

def _add_heading(doc, text, size=16):
    p = doc.add_paragraph()
    _rtl_run(p.add_run(text), bold=True, size=size)
    _rtl_para(p)

def _add_para(doc, text, size=12):
    p = doc.add_paragraph()
    _rtl_run(p.add_run(text), size=size)
    _rtl_para(p)

def _add_table(doc, headers, rows):
    if "م" not in headers:
        headers = ["م"] + list(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _rtl_table(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        _rtl_run(p.add_run(str(h)), bold=True, size=10)
        _rtl_para(p)
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'D9D9D9')
        cell._tc.get_or_add_tcPr().append(shd)
    serial_idx = headers.index("م")
    for idx, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        cells[serial_idx].text = ""
        p = cells[serial_idx].paragraphs[0]
        _rtl_run(p.add_run(str(idx)), size=10)
        _rtl_para(p)
        col = 0
        for i, h in enumerate(headers):
            if i == serial_idx:
                continue
            val = row_data[col] if col < len(row_data) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _rtl_run(p.add_run(str(val) if pd.notna(val) and str(val) != "nan" else ""), size=10)
            _rtl_para(p)
            col += 1

# ============================================================
# ثوابت
# ============================================================
MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]

T_RESEARCH = "خطة بحثية"
T_CONSULT = "استشارية"
T_SELF_FUNDED = "دراسات بتمويل ذاتي"
T_REP_PROGRESS = "تقرير مرحلي"
T_REP_FINAL = "تقرير نهائي"
T_REP_INSPECT = "تقرير معاينة"
T_REP_CALIB = "تقرير معايرة"
T_REP_OTHER = "أخرى"
T_PROP_CONSULT = "استشارات"
T_PROP_FUNDING = "تمويل أبحاث"
T_ACT_CONF = "مؤتمرات"
T_ACT_WORKSHOP = "ورش عمل"
T_MEET_MINISTRY = "بالوزارة"
T_MEET_CENTER = "بالمركز"
T_MEET_STAFF = "مع العاملين"
T_MEET_EXTERNAL = "جهات خارجية"
T_MEET_ONLINE = "اجتماعات افتراضية Online"
T_MEET_OTHER = "جهات أخرى"
T_PAPER_PUBLISHED = "المنشورة"
T_PAPER_ONGOING = "الجاري العمل عليها"
T_THESIS_FINISHED = "المنتهية"
T_THESIS_ONGOING = "الجاري العمل عليها"
T_TRAIN_IN = "داخل المعهد"
T_TRAIN_OUT = "خارج المعهد"
T_TRAIN_ONLINE = "Online"
T_TR_TYPE_COURSE = "دورة تدريبية"
T_TR_TYPE_WS = "ورشة عمل"
T_TR_TYPE_CONF = "مؤتمر"

# ============================================================
# أعمدة KPI الشهرية
# ============================================================
KPI_COLUMNS = [
    "المعهد", "المستخدم", "الشهر", "السنة",
    "خطة بحثية", "استشارية", "دراسات بتمويل ذاتي", "إجمالي الدراسات الجارية",
    "تقرير مرحلي", "تقرير نهائي", "تقرير معاينة", "تقرير معايرة", "أخرى", "إجمالي التقارير",
    "عدد المذكرات",
    "استشارات", "تمويل أبحاث", "إجمالي المقترحات",
    "إجمالي عدد المأموريات", "إجمالي عدد الأيام الفعلية", "إجمالي عدد المعاينات",
    "متدربين", "مدربين", "أيام تدريب داخل المعهد", "أيام تدريب خارج المعهد",
    "مؤتمرات", "ورش عمل", "إجمالي المشاركين في أنشطة علمية",
    "اجتماعات بالوزارة", "اجتماعات بالمركز", "اجتماعات جهات خارجية",
    "اجتماعات مع العاملين", "اجتماعات جهات أخرى", "اجتماعات افتراضية Online", "إجمالي الاجتماعات",
    "الأوراق المنشورة", "الأوراق الجاري العمل عليها", "إجمالي الأوراق البحثية",
    "الرسائل المنتهية", "الرسائل الجاري العمل عليها", "إجمالي الرسائل",
    "إجمالي عدد العاملين الفعلي",
    "الأساتذة المتفرغين", "الأساتذة", "الأساتذة المساعدون",
    "الباحثون", "الباحث المساعد", "مساعد باحث",
    "الفنيون", "الإداريون", "السائقون",
]

# ============================================================
# ⭐ أعمدة المدخلات الربع سنوية الرئيسية (أرقام)
# ============================================================
QUARTERLY_KPI_COLUMNS = [
    "المعهد", "المستخدم", "الربع", "السنة المالية",
    "عدد المشاريع الفردية", "نسبة تنفيذ المشاريع الفردية",
    "عدد المشاريع المشتركة", "نسبة تنفيذ المشاريع المشتركة", "نسبة المشاركة المشتركة",
    "عدد المشاريع الممولة محلياً", "عدد المشاريع الممولة دولياً",
    "عدد التكليفات من الوزارة", "عدد الاستشارات من الوزارة", "عدد الاستشارات من جهات حكومية/خاصة",
    "عدد المقترحات محلياً", "عدد المقترحات دولياً",
    "عدد المقترحات للوزارة", "عدد المقترحات خارج الوزارة",
    "عدد تعاون مع المعاهد", "عدد تعاون مع جهات خارج المعاهد", "عدد تعاون خارجي",
    "عدد الأبحاث المنشورة", "عدد الأبحاث Q1/Q2", "عدد الأبحاث بموافقة نشر",
    "عدد الأبحاث بمكافأة نشر", "عدد الباحثين بمكافأة",
    "عدد رسائل الماجستير", "عدد رسائل الدكتوراه",
    "مراجعة وتحكيم بمجلات", "إشراف لباحثي المعهد", "إشراف لباحثين خارجيين",
    "عدد العضويات الدولية", "عدد العضويات المحلية",
    "دورات بالداخل عدد", "دورات بالداخل متدربين",
    "دورات بالخارج عدد", "دورات بالخارج متدربين",
    "دورات online عدد", "دورات online متدربين",
    "ورش داخل عدد", "ورش داخل متدربين",
    "ورش خارج عدد", "ورش خارج متدربين",
    "مؤتمرات داخل عدد", "مؤتمرات داخل متدربين",
    "مؤتمرات خارج عدد", "مؤتمرات خارج متدربين",
    "دورات نظمها المعهد عدد", "دورات نظمها المعهد متدربين",
    "تمويل الخطة مبلغ", "تمويل الخطة نسبة",
    "تمويل خارجي مبلغ", "تمويل خارجي نسبة",
    "منصرف فعلي مبلغ", "منصرف فعلي نسبة",
    "تمويل نشر دراسات مبلغ", "تمويل نشر أبحاث مبلغ", "تمويل نشر مكافآت مبلغ",
    "عدد الاجتماعات الدورية", "عدد اجتماعات مجلس الإدارة", "عدد اجتماعات وحدة الفكر",
    "آليات التسويق", "آليات الثواب والعقاب",
    "ملاحظات عامة",
]

HALF_COLUMNS = [
    "المعهد", "المستخدم", "النصف", "السنة المالية",
    "إجمالي عدد دراسات الخطة البحثية المنجزة",
    "إجمالي عدد المخرجات التطبيقية من دراسات ومشروعات",
    "إجمالي مبلغ الإيرادات الذاتية",
]

ANNUAL_COLUMNS = ["المعهد", "المستخدم", "السنة المالية", "عدد براءات الاختراع"]

# ============================================================
# القوائم الشهرية
# ============================================================
LIST_STUDIES = ["م", "عنوان الدراسة", "النوع", "الجهة المستفيدة", "تاريخ اصدار التقرير", "ملخص التقرير"]
LIST_REPORTS = ["م", "عنوان التقرير", "النوع", "تاريخ الاصدار", "الجهة المستفيدة", "عنوان الدراسة طبقاً للخطة البحثية", "ملخص"]
LIST_MEMOS = ["م", "رقم المذكرة", "الموضوع", "التاريخ", "الجهة الموجه إليها"]
LIST_PROPOSALS = ["م", "عنوان المقترح", "النوع", "الجهة التمويل", "موقف المقترح", "ملاحظات"]
LIST_MISSIONS = ["م", "الموقع", "الغرض", "من تاريخ", "إلى تاريخ", "عدد الأيام", "الفريق", "هل معاينة؟"]
LIST_TRAINING = ["م", "عنوان البرنامج", "النوع", "المكان", "عدد المشاركين", "مدة البرنامج"]
LIST_ACTIVITIES = ["م", "النشاط", "النوع", "عدد المشاركين", "المكان", "التاريخ"]
LIST_MEETINGS = ["م", "التاريخ", "النوع", "الاسم", "الوظيفة", "مكان انعقاد الاجتماع", "الغرض من الاجتماع"]
LIST_PAPERS = ["م", "عنوان البحث", "المجلة", "النوع", "تاريخ النشر", "الباحث الرئيسي"]
LIST_THESES = ["م", "عنوان الرسالة", "النوع", "الباحث", "الجامعة", "تاريخ المناقشة"]
LIST_PERF_MEETINGS = ["م", "تاريخ الاجتماع", "عنوان الاجتماع", "الغرض من الاجتماع", "السادة الحضور"]

MONTHLY_LISTS = {
    "الدراسات_الجارية": {"title": "📚 الدراسات الجارية", "cols": LIST_STUDIES,
                         "hint": f"النوع: {T_RESEARCH} / {T_CONSULT} / {T_SELF_FUNDED}",
                         "dropdowns": {"النوع": [T_RESEARCH, T_CONSULT, T_SELF_FUNDED]}},
    "التقارير_الصادرة": {"title": "📄 التقارير الصادرة", "cols": LIST_REPORTS,
                         "hint": f"النوع: {T_REP_PROGRESS} / {T_REP_FINAL} / {T_REP_INSPECT} / {T_REP_CALIB} / {T_REP_OTHER}",
                         "dropdowns": {"النوع": [T_REP_PROGRESS, T_REP_FINAL, T_REP_INSPECT, T_REP_CALIB, T_REP_OTHER]}},
    "المذكرات_الصادرة": {"title": "📝 المذكرات الصادرة", "cols": LIST_MEMOS, "hint": "", "dropdowns": {}},
    "المقترحات": {"title": "📋 المقترحات الفنية والمالية", "cols": LIST_PROPOSALS,
                  "hint": f"النوع: {T_PROP_CONSULT} / {T_PROP_FUNDING}",
                  "dropdowns": {"النوع": [T_PROP_CONSULT, T_PROP_FUNDING]}},
    "المأموريات_الحقلية": {"title": "🏃 المأموريات الحقلية والمعاينات", "cols": LIST_MISSIONS,
                            "hint": "أدخل عدد الأيام لكل مأمورية. عمود (هل معاينة؟) يشير إلى نوع الزيارة",
                            "dropdowns": {"هل معاينة؟": ["مأمورية", "معاينة"]}},
    "التدريب": {"title": "👥 التدريب", "cols": LIST_TRAINING,
                "hint": f"النوع: {T_TR_TYPE_COURSE} / {T_TR_TYPE_WS} / {T_TR_TYPE_CONF} — المكان: {T_TRAIN_IN} / {T_TRAIN_OUT} / {T_TRAIN_ONLINE}",
                "dropdowns": {"النوع": [T_TR_TYPE_COURSE, T_TR_TYPE_WS, T_TR_TYPE_CONF],
                              "المكان": [T_TRAIN_IN, T_TRAIN_OUT, T_TRAIN_ONLINE]}},
    "الأنشطة_العلمية": {"title": "🔬 الأنشطة العلمية", "cols": LIST_ACTIVITIES,
                        "hint": f"النوع: {T_ACT_CONF} / {T_ACT_WORKSHOP}",
                        "dropdowns": {"النوع": [T_ACT_CONF, T_ACT_WORKSHOP]}},
    "الاجتماعات": {"title": "🤝 الاجتماعات", "cols": LIST_MEETINGS,
                   "hint": f"النوع: {T_MEET_MINISTRY} / {T_MEET_CENTER} / {T_MEET_STAFF} / {T_MEET_EXTERNAL} / {T_MEET_ONLINE} / {T_MEET_OTHER}",
                   "dropdowns": {"النوع": [T_MEET_MINISTRY, T_MEET_CENTER, T_MEET_STAFF, T_MEET_EXTERNAL, T_MEET_ONLINE, T_MEET_OTHER]}},
    "الأوراق_البحثية": {"title": "📝 الأوراق البحثية", "cols": LIST_PAPERS,
                        "hint": f"النوع: {T_PAPER_PUBLISHED} / {T_PAPER_ONGOING}",
                        "dropdowns": {"النوع": [T_PAPER_PUBLISHED, T_PAPER_ONGOING]}},
    "الرسائل_العلمية": {"title": "🎓 رسائل الماجستير والدكتوراه", "cols": LIST_THESES,
                        "hint": f"النوع: {T_THESIS_FINISHED} / {T_THESIS_ONGOING}",
                        "dropdowns": {"النوع": [T_THESIS_FINISHED, T_THESIS_ONGOING]}},
    "اجتماعات_تحسين_الأداء": {"title": "👥 اجتماعات تحسين الأداء", "cols": LIST_PERF_MEETINGS, "hint": "", "dropdowns": {}},
}

# ============================================================
# ⭐ القوائم التفصيلية الربع سنوية
# ============================================================
Q_PROJECTS_INDIV = ["م", "عنوان المشروع البحثي", "نسبة المنفذ %", "عدد التقارير", "هدف الدراسة", "النتائج"]
Q_PROJECTS_JOINT = ["م", "عنوان المشروع البحثي", "نسبة المنفذ %", "نسبة المشاركة %", "عدد التقارير", "الهدف", "النتائج"]
Q_PROJECTS_LOCAL = ["م", "عنوان المشروع", "الجهة الممولة", "الأهداف"]
Q_PROJECTS_INTL = ["م", "عنوان المشروع", "الجهة الخارجية", "موقف المقترح"]
Q_MINISTRY_ASSIGNMENTS = ["م", "عنوان الدراسة", "القطاع/الهيئة", "الموقف", "المدة"]
Q_MINISTRY_CONSULT = ["م", "عنوان الدراسة", "الجهة المكلفة", "الموقف"]
Q_OTHER_CONSULT = ["م", "عنوان الدراسة", "الجهة المكلفة", "الموقف"]
Q_PROPOSALS_LOCAL = ["م", "عنوان المقترح", "الجهة التمويل", "موقف المقترح"]
Q_PROPOSALS_INTL = ["م", "عنوان المقترح", "الجهة التمويل", "موقف المقترح"]
Q_PROPOSALS_MINISTRY = ["م", "عنوان المقترح", "القطاع/الهيئة", "موقف المقترح"]
Q_PROPOSALS_OUTSIDE = ["م", "عنوان المقترح", "الوزارة/الهيئة", "موقف المقترح"]
Q_COLLAB_INSTITUTES = ["م", "اسم الدراسة/المشروع/بروتوكول", "المعاهد", "تاريخ بدء التعاون"]
Q_COLLAB_ENTITIES = ["م", "اسم الدراسة/المشروع/بروتوكول", "الجهات", "تاريخ بدء التعاون"]
Q_COLLAB_FOREIGN = ["م", "اسم الدراسة/المشروع/بروتوكول", "الجهات", "تاريخ بدء التعاون"]
Q_INTL_MEMBERSHIPS = ["م", "المنظمة", "نوع العضوية"]
Q_LOCAL_MEMBERSHIPS = ["م", "المنظمة", "نوع العضوية"]
Q_MINISTRY_MEETINGS = ["م", "التاريخ", "الاسم", "الوظيفة", "مكان انعقاد الاجتماع", "الغرض من الاجتماع"]

QUARTERLY_LISTS = {
    "Q_مشاريع_فردية": {"title": "1-1-1 مشاريع فردية", "cols": Q_PROJECTS_INDIV},
    "Q_مشاريع_مشتركة": {"title": "1-1-2 مشاريع مشتركة", "cols": Q_PROJECTS_JOINT},
    "Q_مشاريع_ممولة_محلياً": {"title": "2-1-1 مشاريع ممولة محلياً", "cols": Q_PROJECTS_LOCAL},
    "Q_مشاريع_ممولة_دولياً": {"title": "2-1-2 مشاريع ممولة دولياً", "cols": Q_PROJECTS_INTL},
    "Q_تكليفات_وزارة": {"title": "2-1-3 تكليفات الوزارة", "cols": Q_MINISTRY_ASSIGNMENTS},
    "Q_استشارات_وزارة": {"title": "2-1-4 استشارات الوزارة", "cols": Q_MINISTRY_CONSULT},
    "Q_استشارات_جهات_أخرى": {"title": "2-1-5 استشارات جهات أخرى", "cols": Q_OTHER_CONSULT},
    "Q_مقترحات_محلية": {"title": "2-2-1 مقترحات محلية", "cols": Q_PROPOSALS_LOCAL},
    "Q_مقترحات_دولية": {"title": "2-2-2 مقترحات دولية", "cols": Q_PROPOSALS_INTL},
    "Q_مقترحات_وزارة": {"title": "2-2-3 مقترحات للوزارة", "cols": Q_PROPOSALS_MINISTRY},
    "Q_مقترحات_خارج_وزارة": {"title": "2-2-4 مقترحات خارج الوزارة", "cols": Q_PROPOSALS_OUTSIDE},
    "Q_تعاون_معاهد": {"title": "3-1 تعاون مع المعاهد", "cols": Q_COLLAB_INSTITUTES},
    "Q_تعاون_جهات": {"title": "3-2 تعاون مع جهات داخلية", "cols": Q_COLLAB_ENTITIES},
    "Q_تعاون_خارجي": {"title": "3-3 تعاون خارجي", "cols": Q_COLLAB_FOREIGN},
    "Q_عضويات_دولية": {"title": "4-2 عضويات دولية", "cols": Q_INTL_MEMBERSHIPS},
    "Q_عضويات_محلية": {"title": "4-2 عضويات محلية", "cols": Q_LOCAL_MEMBERSHIPS},
    "Q_اجتماعات_وزارية": {"title": "8 الاجتماعات الوزارية", "cols": Q_MINISTRY_MEETINGS},
}

# ============================================================
# الفترة المالية
# ============================================================
def get_fiscal_period(month, year):
    result = {
        "fiscal_year": "", "quarter": None, "half": None,
        "months_for_quarter": [], "months_for_half": [],
        "is_quarter_month": month in ["أكتوبر", "يناير", "أبريل", "يوليو"],
        "is_half_month": month in ["يناير", "يوليو"],
        "is_annual_month": month == "يوليو",
    }
    if month == "أكتوبر":
        result["quarter"] = "الربع الأول"
        result["months_for_quarter"] = [("يوليو", year), ("أغسطس", year), ("سبتمبر", year)]
        result["fiscal_year"] = f"{year}/{year + 1}"
    elif month == "يناير":
        result["quarter"] = "الربع الثاني"
        result["months_for_quarter"] = [("أكتوبر", year - 1), ("نوفمبر", year - 1), ("ديسمبر", year - 1)]
        result["fiscal_year"] = f"{year - 1}/{year}"
    elif month == "أبريل":
        result["quarter"] = "الربع الثالث"
        result["months_for_quarter"] = [("يناير", year), ("فبراير", year), ("مارس", year)]
        result["fiscal_year"] = f"{year - 1}/{year}"
    elif month == "يوليو":
        result["quarter"] = "الربع الرابع"
        result["months_for_quarter"] = [("أبريل", year), ("مايو", year), ("يونيو", year)]
        result["fiscal_year"] = f"{year - 1}/{year}"
    else:
        month_idx = MONTHS.index(month) + 1
        result["fiscal_year"] = f"{year}/{year + 1}" if month_idx >= 7 else f"{year - 1}/{year}"
    if month == "يناير":
        result["half"] = "النصف الأول"
    elif month == "يوليو":
        result["half"] = "النصف الثاني"
    return result

def get_quarter_months(quarter, fiscal_year):
    fy_s = int(fiscal_year.split("/")[0]); fy_e = int(fiscal_year.split("/")[1])
    if quarter == "الربع الأول":
        return [("يوليو", fy_s), ("أغسطس", fy_s), ("سبتمبر", fy_s)]
    elif quarter == "الربع الثاني":
        return [("أكتوبر", fy_s), ("نوفمبر", fy_s), ("ديسمبر", fy_s)]
    elif quarter == "الربع الثالث":
        return [("يناير", fy_e), ("فبراير", fy_e), ("مارس", fy_e)]
    return [("أبريل", fy_e), ("مايو", fy_e), ("يونيو", fy_e)]

# ============================================================
# حساب KPI الشهري
# ============================================================
def compute_kpi_from_lists(institute, month, year):
    period = f"{month} {year}"
    kpi = {c: 0 for c in KPI_COLUMNS if c not in ["المعهد", "المستخدم", "الشهر", "السنة"]}
    
    def get_list(name):
        save_key = f"_saved_{name}_{institute}_{period}"
        if save_key in st.session_state:
            df = st.session_state[save_key].copy()
            df["المعهد"] = institute; df["الفترة"] = period
            return df
        cols = MONTHLY_LISTS[name]["cols"] + ["المعهد", "الفترة"]
        df = load_sheet(name, cols)
        if df.empty or "المعهد" not in df.columns or "الفترة" not in df.columns:
            return pd.DataFrame()
        return df[(df["المعهد"] == institute) & (df["الفترة"] == period)]
    
    def to_num(series):
        return pd.to_numeric(series, errors="coerce").fillna(0)
    
    s = get_list("الدراسات_الجارية")
    if not s.empty and "النوع" in s.columns:
        kpi["خطة بحثية"] = int((s["النوع"] == T_RESEARCH).sum())
        kpi["استشارية"] = int((s["النوع"] == T_CONSULT).sum())
        kpi["دراسات بتمويل ذاتي"] = int((s["النوع"] == T_SELF_FUNDED).sum())
    kpi["إجمالي الدراسات الجارية"] = kpi["خطة بحثية"] + kpi["استشارية"] + kpi["دراسات بتمويل ذاتي"]
    
    s = get_list("التقارير_الصادرة")
    if not s.empty and "النوع" in s.columns:
        kpi["تقرير مرحلي"] = int((s["النوع"] == T_REP_PROGRESS).sum())
        kpi["تقرير نهائي"] = int((s["النوع"] == T_REP_FINAL).sum())
        kpi["تقرير معاينة"] = int((s["النوع"] == T_REP_INSPECT).sum())
        kpi["تقرير معايرة"] = int((s["النوع"] == T_REP_CALIB).sum())
        kpi["أخرى"] = int((s["النوع"] == T_REP_OTHER).sum())
    kpi["إجمالي التقارير"] = kpi["تقرير مرحلي"] + kpi["تقرير نهائي"] + kpi["تقرير معاينة"] + kpi["تقرير معايرة"] + kpi["أخرى"]
    
    s = get_list("المذكرات_الصادرة")
    kpi["عدد المذكرات"] = len(s) if not s.empty else 0
    
    s = get_list("المقترحات")
    if not s.empty and "النوع" in s.columns:
        kpi["استشارات"] = int((s["النوع"] == T_PROP_CONSULT).sum())
        kpi["تمويل أبحاث"] = int((s["النوع"] == T_PROP_FUNDING).sum())
    kpi["إجمالي المقترحات"] = kpi["استشارات"] + kpi["تمويل أبحاث"]
    
    s = get_list("المأموريات_الحقلية")
    if not s.empty:
        if "عدد الأيام" in s.columns:
            kpi["إجمالي عدد الأيام الفعلية"] = int(to_num(s["عدد الأيام"]).sum())
        if "هل معاينة؟" in s.columns:
            kpi["إجمالي عدد المأموريات"] = int((s["هل معاينة؟"] == "مأمورية").sum())
            kpi["إجمالي عدد المعاينات"] = int((s["هل معاينة؟"] == "معاينة").sum())
    
    s = get_list("التدريب")
    if not s.empty:
        if "عدد المشاركين" in s.columns:
            kpi["متدربين"] = int(to_num(s["عدد المشاركين"]).sum())
        if "المكان" in s.columns and "مدة البرنامج" in s.columns:
            days = to_num(s["مدة البرنامج"])
            kpi["أيام تدريب داخل المعهد"] = int(days[s["المكان"] == T_TRAIN_IN].sum())
            kpi["أيام تدريب خارج المعهد"] = int(days[s["المكان"] == T_TRAIN_OUT].sum())
    
    s = get_list("الأنشطة_العلمية")
    if not s.empty and "عدد المشاركين" in s.columns and "النوع" in s.columns:
        parts = to_num(s["عدد المشاركين"])
        kpi["مؤتمرات"] = int(parts[s["النوع"] == T_ACT_CONF].sum())
        kpi["ورش عمل"] = int(parts[s["النوع"] == T_ACT_WORKSHOP].sum())
    kpi["إجمالي المشاركين في أنشطة علمية"] = kpi["مؤتمرات"] + kpi["ورش عمل"]
    
    s = get_list("الاجتماعات")
    if not s.empty and "النوع" in s.columns:
        kpi["اجتماعات بالوزارة"] = int((s["النوع"] == T_MEET_MINISTRY).sum())
        kpi["اجتماعات بالمركز"] = int((s["النوع"] == T_MEET_CENTER).sum())
        kpi["اجتماعات جهات خارجية"] = int((s["النوع"] == T_MEET_EXTERNAL).sum())
        kpi["اجتماعات مع العاملين"] = int((s["النوع"] == T_MEET_STAFF).sum())
        kpi["اجتماعات جهات أخرى"] = int((s["النوع"] == T_MEET_OTHER).sum())
        kpi["اجتماعات افتراضية Online"] = int((s["النوع"] == T_MEET_ONLINE).sum())
    kpi["إجمالي الاجتماعات"] = sum([kpi["اجتماعات بالوزارة"], kpi["اجتماعات بالمركز"], kpi["اجتماعات جهات خارجية"],
                                     kpi["اجتماعات مع العاملين"], kpi["اجتماعات جهات أخرى"], kpi["اجتماعات افتراضية Online"]])
    
    s = get_list("الأوراق_البحثية")
    if not s.empty and "النوع" in s.columns:
        kpi["الأوراق المنشورة"] = int((s["النوع"] == T_PAPER_PUBLISHED).sum())
        kpi["الأوراق الجاري العمل عليها"] = int((s["النوع"] == T_PAPER_ONGOING).sum())
    kpi["إجمالي الأوراق البحثية"] = kpi["الأوراق المنشورة"] + kpi["الأوراق الجاري العمل عليها"]
    
    s = get_list("الرسائل_العلمية")
    if not s.empty and "النوع" in s.columns:
        kpi["الرسائل المنتهية"] = int((s["النوع"] == T_THESIS_FINISHED).sum())
        kpi["الرسائل الجاري العمل عليها"] = int((s["النوع"] == T_THESIS_ONGOING).sum())
    kpi["إجمالي الرسائل"] = kpi["الرسائل المنتهية"] + kpi["الرسائل الجاري العمل عليها"]
    return kpi

# ============================================================
# ⭐ التقرير الشهري
# ============================================================
def generate_monthly_report(institute, month, year):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = Cm(2); section.right_margin = Cm(2)
    
    doc.add_paragraph(); doc.add_paragraph()
    _add_heading(doc, "تقرير الإنجازات الشهري", size=26)
    _add_heading(doc, f"{institute}", size=20)
    _add_heading(doc, f"({month} {year})", size=18)
    doc.add_page_break()
    
    _add_heading(doc, "المحتويات", size=20)
    for i, c in enumerate(["الدراسات البحثية", "التقارير الفنية", "الاجتماعات", "التدريب والندوات", "اجتماعات تحسين الأداء"], 1):
        _add_para(doc, f"{i}. {c}", size=13)
    doc.add_page_break()
    
    period = f"{month} {year}"
    kpi = compute_kpi_from_lists(institute, month, year)
    
    _add_heading(doc, "ملخص الإنجازات الشهرية", size=16)
    summary = [
        ("إجمالي الدراسات الجارية", kpi["إجمالي الدراسات الجارية"]),
        ("إجمالي التقارير الصادرة", kpi["إجمالي التقارير"]),
        ("عدد المذكرات الصادرة", kpi["عدد المذكرات"]),
        ("إجمالي المقترحات", kpi["إجمالي المقترحات"]),
        ("إجمالي المأموريات", kpi["إجمالي عدد المأموريات"]),
        ("إجمالي أيام المأموريات", kpi["إجمالي عدد الأيام الفعلية"]),
        ("إجمالي المعاينات", kpi["إجمالي عدد المعاينات"]),
        ("عدد المتدربين", kpi["متدربين"]),
        ("إجمالي المشاركين في أنشطة علمية", kpi["إجمالي المشاركين في أنشطة علمية"]),
        ("إجمالي الاجتماعات", kpi["إجمالي الاجتماعات"]),
        ("إجمالي الأوراق البحثية", kpi["إجمالي الأوراق البحثية"]),
        ("إجمالي الرسائل العلمية", kpi["إجمالي الرسائل"]),
    ]
    _add_table(doc, ["البند", "العدد"], [[label, str(val)] for label, val in summary])
    doc.add_page_break()
    
    # الدراسات
    _add_heading(doc, f"الدراسات البحثية - {month} {year}", size=18)
    df_studies = load_sheet("الدراسات_الجارية", LIST_STUDIES + ["المعهد", "الفترة"])
    df_filtered = pd.DataFrame()
    if not df_studies.empty:
        df_filtered = df_studies[(df_studies["المعهد"] == institute) & (df_studies["الفترة"] == period)]
    
    for title, tval in [("دراسات الخطة البحثية", T_RESEARCH), ("مشروعات بحثية ممولة", T_SELF_FUNDED), ("استشارات بحثية ممولة", T_CONSULT)]:
        _add_heading(doc, title, size=14)
        rows = []
        if not df_filtered.empty and "النوع" in df_filtered.columns:
            for _, row in df_filtered[df_filtered["النوع"] == tval].iterrows():
                rows.append([row.get("عنوان الدراسة",""), row.get("الجهة المستفيدة",""), row.get("تاريخ اصدار التقرير",""), row.get("ملخص التقرير","")])
        _add_table(doc, ["عنوان الدراسة", "الجهة المستفيدة", "تاريخ اصدار التقرير", "ملخص التقرير"], rows)
    doc.add_page_break()
    
    # التقارير
    _add_heading(doc, "التقارير الفنية", size=18)
    df_reports = load_sheet("التقارير_الصادرة", LIST_REPORTS + ["المعهد", "الفترة"])
    rows = []
    if not df_reports.empty:
        r = df_reports[(df_reports["المعهد"] == institute) & (df_reports["الفترة"] == period)]
        for _, row in r.iterrows():
            rows.append([row.get("عنوان التقرير",""), row.get("تاريخ الاصدار",""), row.get("الجهة المستفيدة",""),
                        row.get("عنوان الدراسة طبقاً للخطة البحثية",""), row.get("ملخص","")])
    _add_table(doc, ["عنوان التقرير", "تاريخ الاصدار", "الجهة المستفيدة", "عنوان الدراسة طبقاً للخطة البحثية", "ملخص"], rows)
    doc.add_page_break()
    
    # الاجتماعات
    _add_heading(doc, "الاجتماعات", size=18)
    df_meetings = load_sheet("الاجتماعات", LIST_MEETINGS + ["المعهد", "الفترة"])
    df_meet = pd.DataFrame()
    if not df_meetings.empty:
        df_meet = df_meetings[(df_meetings["المعهد"] == institute) & (df_meetings["الفترة"] == period)]
    
    for title, tval in [("أولاً: الاجتماعات بالوزارة", T_MEET_MINISTRY), ("ثانياً: اجتماعات ديوان عام المركز", T_MEET_CENTER),
                         ("ثالثاً: الاجتماعات بمقر المعهد", T_MEET_STAFF), ("رابعاً: اجتماعات بجهات أخرى", T_MEET_EXTERNAL),
                         ("خامساً: اجتماعات افتراضية Online", T_MEET_ONLINE)]:
        _add_heading(doc, title, size=14)
        rows = []
        if not df_meet.empty and "النوع" in df_meet.columns:
            for _, row in df_meet[df_meet["النوع"] == tval].iterrows():
                rows.append([row.get("التاريخ",""), row.get("الاسم",""), row.get("الوظيفة",""),
                            row.get("مكان انعقاد الاجتماع",""), row.get("الغرض من الاجتماع","")])
        _add_table(doc, ["التاريخ", "الاسم", "الوظيفة", "مكان انعقاد الاجتماع", "الغرض من الاجتماع"], rows)
    doc.add_page_break()
    
    # التدريب
    _add_heading(doc, "الدورات التدريبية / ورش العمل / الندوات / المؤتمرات", size=18)
    df_training = load_sheet("التدريب", LIST_TRAINING + ["المعهد", "الفترة"])
    rows = []
    if not df_training.empty:
        r = df_training[(df_training["المعهد"] == institute) & (df_training["الفترة"] == period)]
        for _, row in r.iterrows():
            rows.append([row.get("عنوان البرنامج",""), row.get("عدد المشاركين",""), row.get("المكان",""), row.get("مدة البرنامج","")])
    _add_table(doc, ["عنوان البرنامج", "عدد المشاركين", "مكان الانعقاد", "مدة البرنامج"], rows)
    doc.add_page_break()
    
    # اجتماعات تحسين الأداء
    _add_heading(doc, "اجتماعات تحسين الأداء وتعزيز الروح الجماعية", size=18)
    df_perf = load_sheet("اجتماعات_تحسين_الأداء", LIST_PERF_MEETINGS + ["المعهد", "الفترة"])
    rows = []
    if not df_perf.empty:
        r = df_perf[(df_perf["المعهد"] == institute) & (df_perf["الفترة"] == period)]
        for _, row in r.iterrows():
            rows.append([row.get("تاريخ الاجتماع",""), row.get("عنوان الاجتماع",""), row.get("الغرض من الاجتماع",""), row.get("السادة الحضور","")])
    _add_table(doc, ["تاريخ الاجتماع", "عنوان الاجتماع", "الغرض من الاجتماع", "السادة الحضور"], rows)
    
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()

# ============================================================
# ⭐ التقرير الربع سنوي (شامل 8 أقسام)
# ============================================================
def generate_quarterly_report(institute, quarter, fiscal_year):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    
    doc.add_paragraph(); doc.add_paragraph()
    _add_heading(doc, "تقرير مؤشرات الأداء الربع سنوي", size=26)
    _add_heading(doc, f"{institute}", size=20)
    _add_heading(doc, f"({quarter} — السنة المالية {fiscal_year})", size=16)
    doc.add_page_break()
    
    # تحميل البيانات الربع سنوية
    df_qk = load_sheet("Quarterly_KPI", QUARTERLY_KPI_COLUMNS)
    q_row = None
    if not df_qk.empty:
        r = df_qk[(df_qk["المعهد"] == institute) & (df_qk["الربع"] == quarter) & (df_qk["السنة المالية"] == fiscal_year)]
        if not r.empty:
            q_row = r.iloc[0]
    
    def qval(col, default=0):
        if q_row is None or col not in q_row.index: return default
        try:
            v = q_row[col]
            if pd.isna(v) or v == "": return default
            return int(float(v))
        except: return default
    
    def qvalf(col, default=0.0):
        if q_row is None or col not in q_row.index: return default
        try:
            v = q_row[col]
            if pd.isna(v) or v == "": return default
            return float(v)
        except: return default
    
    def qstr(col, default=""):
        if q_row is None or col not in q_row.index: return default
        return str(q_row[col]) if pd.notna(q_row[col]) else default
    
    def load_q_list(sheet_name):
        cols = QUARTERLY_LISTS[sheet_name]["cols"] + ["المعهد", "الفترة"]
        df = load_sheet(sheet_name, cols)
        if df.empty or "المعهد" not in df.columns or "الفترة" not in df.columns:
            return pd.DataFrame()
        return df[(df["المعهد"] == institute) & (df["الفترة"] == f"{quarter} - {fiscal_year}")]
    
    period_q = f"{quarter} - {fiscal_year}"
    
    # ============ 1. الأداء التنفيذي للخطة البحثية ============
    _add_heading(doc, "1. الأداء التنفيذي للخطة البحثية (30%)", size=16)
    _add_para(doc, "تعتبر الخطة البحثية أحد الركائز الأساسية لعمل المعهد. تتضمن هذه الفقرة متابعة تنفيذ المشروعات البحثية وفقاً للخطة الموضوعة، مع تقييم نسبة الإنجاز وإصدار التقارير الدورية.")
    
    _add_heading(doc, "1-1 عدد مشروعات الخطة البحثية ونسبة المنفذ من كل مشروع والتقارير الصادرة:", size=13)
    
    _add_heading(doc, "1-1-1 مشروع فردي خاص بالمعهد فقط", size=13)
    df_ind = load_q_list("Q_مشاريع_فردية")
    rows = []
    if not df_ind.empty:
        for _, row in df_ind.iterrows():
            rows.append([row.get("عنوان المشروع البحثي",""), row.get("نسبة المنفذ %",""),
                        row.get("عدد التقارير",""), row.get("هدف الدراسة",""), row.get("النتائج","")])
    _add_table(doc, ["عنوان المشروع البحثي", "نسبة المنفذ %", "عدد التقارير", "هدف الدراسة", "النتائج"], rows)
    _add_para(doc, f"التقييم العام: عدد {qval('عدد المشاريع الفردية')} مشروع، نسبة تنفيذ {qvalf('نسبة تنفيذ المشاريع الفردية'):.1f}%", size=11)
    
    _add_heading(doc, "1-1-2 مشروع مشترك مع معاهد أخرى", size=13)
    df_jnt = load_q_list("Q_مشاريع_مشتركة")
    rows = []
    if not df_jnt.empty:
        for _, row in df_jnt.iterrows():
            rows.append([row.get("عنوان المشروع البحثي",""), row.get("نسبة المنفذ %",""), row.get("نسبة المشاركة %",""),
                        row.get("عدد التقارير",""), row.get("الهدف",""), row.get("النتائج","")])
    _add_table(doc, ["عنوان المشروع البحثي", "نسبة المنفذ %", "نسبة المشاركة", "عدد التقارير", "الهدف", "النتائج"], rows)
    _add_para(doc, f"التقييم العام: عدد {qval('عدد المشاريع المشتركة')} مشروع، نسبة تنفيذ {qvalf('نسبة تنفيذ المشاريع المشتركة'):.1f}%", size=11)
    doc.add_page_break()
    
    # ============ 2. الاستشارات والدراسات خارج الخطة ============
    _add_heading(doc, "2. الاستشارات والدراسات خارج الخطة البحثية (15%)", size=16)
    _add_para(doc, "يتعلق هذا الجزء بالدراسات والاستشارات التي تمت خارج إطار الخطة البحثية الأساسية للمعهد.")
    
    _add_heading(doc, "2-1 عدد المشاريع التي تمت خلال الفترة", size=13)
    total_2_1 = (qval("عدد المشاريع الممولة محلياً") + qval("عدد المشاريع الممولة دولياً")
                 + qval("عدد التكليفات من الوزارة") + qval("عدد الاستشارات من الوزارة")
                 + qval("عدد الاستشارات من جهات حكومية/خاصة"))
    rows = [
        ["عدد المشاريع البحثية الممولة محلياً (STDF)", str(qval("عدد المشاريع الممولة محلياً"))],
        ["عدد المشاريع البحثية الممولة دولياً (PRIMA)", str(qval("عدد المشاريع الممولة دولياً"))],
        ["عدد التكليفات من الوزارة", str(qval("عدد التكليفات من الوزارة"))],
        ["عدد الاستشارات من الوزارة", str(qval("عدد الاستشارات من الوزارة"))],
        ["عدد الاستشارات من جهات حكومية/خاصة", str(qval("عدد الاستشارات من جهات حكومية/خاصة"))],
        ["الإجمالي", str(total_2_1)],
    ]
    _add_table(doc, ["البند", "العدد"], rows)
    
    # 2-1-1 مشاريع ممولة محلياً
    _add_heading(doc, "2-1-1 المشاريع البحثية الممولة محلياً", size=13)
    df = load_q_list("Q_مشاريع_ممولة_محلياً")
    rows = [[r.get("عنوان المشروع",""), r.get("الجهة الممولة",""), r.get("الأهداف","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["عنوان المشروع", "الجهة الممولة", "الأهداف"], rows)
    
    _add_heading(doc, "2-1-2 المشاريع البحثية الممولة دولياً (PRIMA)", size=13)
    df = load_q_list("Q_مشاريع_ممولة_دولياً")
    rows = [[r.get("عنوان المشروع",""), r.get("الجهة الخارجية",""), r.get("موقف المقترح","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["عنوان المشروع", "الجهة الخارجية", "موقف المقترح"], rows)
    
    _add_heading(doc, "2-1-3 التكليفات من الوزارة", size=13)
    df = load_q_list("Q_تكليفات_وزارة")
    rows = [[r.get("عنوان الدراسة",""), r.get("القطاع/الهيئة",""), r.get("الموقف",""), r.get("المدة","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["عنوان الدراسة", "القطاع/الهيئة", "الموقف", "المدة"], rows)
    
    _add_heading(doc, "2-1-4 الاستشارات من الوزارة", size=13)
    df = load_q_list("Q_استشارات_وزارة")
    rows = [[r.get("عنوان الدراسة",""), r.get("الجهة المكلفة",""), r.get("الموقف","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["عنوان الدراسة", "الجهة المكلفة", "الموقف"], rows)
    
    _add_heading(doc, "2-1-5 الاستشارات من جهات حكومية/خاصة", size=13)
    df = load_q_list("Q_استشارات_جهات_أخرى")
    rows = [[r.get("عنوان الدراسة",""), r.get("الجهة المكلفة",""), r.get("الموقف","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["عنوان الدراسة", "الجهة المكلفة", "الموقف"], rows)
    doc.add_page_break()
    
    # 2-2 المقترحات
    _add_heading(doc, "2-2 المقترحات البحثية في إعداد التقديم (Pipeline)", size=13)
    total_2_2 = (qval("عدد المقترحات محلياً") + qval("عدد المقترحات دولياً")
                 + qval("عدد المقترحات للوزارة") + qval("عدد المقترحات خارج الوزارة"))
    rows = [
        ["عدد المقترحات المقدمة محلياً", str(qval("عدد المقترحات محلياً"))],
        ["عدد المقترحات المقدمة دولياً", str(qval("عدد المقترحات دولياً"))],
        ["عدد المقترحات المقدمة للوزارة", str(qval("عدد المقترحات للوزارة"))],
        ["عدد المقترحات لجهات خارج الوزارة", str(qval("عدد المقترحات خارج الوزارة"))],
        ["الإجمالي", str(total_2_2)],
    ]
    _add_table(doc, ["البند", "العدد"], rows)
    
    for title, sheet in [("2-2-1 المقترحات المقدمة محلياً", "Q_مقترحات_محلية"),
                          ("2-2-2 المقترحات المقدمة دولياً", "Q_مقترحات_دولية"),
                          ("2-2-3 المقترحات المقدمة للوزارة", "Q_مقترحات_وزارة"),
                          ("2-2-4 المقترحات لجهات خارج الوزارة", "Q_مقترحات_خارج_وزارة")]:
        _add_heading(doc, title, size=13)
        df = load_q_list(sheet)
        rows = [[r.get("عنوان المقترح",""), r.get("الجهة التمويل", r.get("القطاع/الهيئة", r.get("الوزارة/الهيئة",""))), r.get("موقف المقترح","")] for _, r in df.iterrows()] if not df.empty else []
        _add_table(doc, ["عنوان المقترح", "الجهة/الهيئة", "موقف المقترح"], rows)
    doc.add_page_break()
    
    # ============ 3. التعاون الداخلي والخارجي ============
    _add_heading(doc, "3. التعاون الداخلي والخارجي (8%)", size=16)
    _add_para(doc, "يركز هذا الجزء على الشراكات والتعاون مع المعاهد البحثية الأخرى التابعة للمركز أو الجهات والجامعات الحكومية أو الأجنبية.")
    
    _add_heading(doc, "3-1 التعاون الداخلي مع المعاهد البحثية", size=13)
    df = load_q_list("Q_تعاون_معاهد")
    rows = [[r.get("اسم الدراسة/المشروع/بروتوكول",""), r.get("المعاهد",""), r.get("تاريخ بدء التعاون","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["اسم الدراسة/المشروع/البروتوكول", "المعاهد", "تاريخ بدء التعاون"], rows)
    
    _add_heading(doc, "3-2 التعاون الداخلي مع جهات خارج المعاهد", size=13)
    df = load_q_list("Q_تعاون_جهات")
    rows = [[r.get("اسم الدراسة/المشروع/بروتوكول",""), r.get("الجهات",""), r.get("تاريخ بدء التعاون","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["اسم الدراسة/المشروع/البروتوكول", "الجهات", "تاريخ بدء التعاون"], rows)
    
    _add_heading(doc, "3-3 التعاون الخارجي مع جهات خارج مصر", size=13)
    df = load_q_list("Q_تعاون_خارجي")
    rows = [[r.get("اسم الدراسة/المشروع/بروتوكول",""), r.get("الجهات",""), r.get("تاريخ بدء التعاون","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["اسم الدراسة/المشروع/البروتوكول", "الجهات", "تاريخ بدء التعاون"], rows)
    doc.add_page_break()
    
    # ============ 4. الإنتاج العلمي ============
    _add_heading(doc, "4. الإنتاج العلمي والنشر والابتكار (14%)", size=16)
    _add_para(doc, "يعد الإنتاج العلمي والنشر من المؤشرات الحيوية لأداء المعهد.")
    
    _add_heading(doc, "4-1 الإنتاج العلمي والنشر", size=13)
    rows = [
        ["عدد الأبحاث التي تم نشرها لباحثي المعهد", str(qval("عدد الأبحاث المنشورة"))],
        ["عدد الأبحاث المنشورة في مجالات متميزة Q1/Q2", str(qval("عدد الأبحاث Q1/Q2"))],
        ["عدد الأبحاث بأخذ الموافقة على النشر", str(qval("عدد الأبحاث بموافقة نشر"))],
        ["عدد الأبحاث بأخذ مكافأة نشر", str(qval("عدد الأبحاث بمكافأة نشر"))],
        ["عدد الباحثين بمكافأة نشر", str(qval("عدد الباحثين بمكافأة"))],
        ["عدد رسائل الماجستير", str(qval("عدد رسائل الماجستير"))],
        ["عدد رسائل الدكتوراه", str(qval("عدد رسائل الدكتوراه"))],
        ["المشاركة في مراجعة وتحكيم الأبحاث بمجلات علمية", str(qval("مراجعة وتحكيم بمجلات"))],
        ["عدد الرسائل المشرف عليها لباحثي المعهد", str(qval("إشراف لباحثي المعهد"))],
        ["عدد الرسائل المشرف عليها لباحثين خارج المعهد", str(qval("إشراف لباحثين خارجيين"))],
    ]
    _add_table(doc, ["البند", "العدد"], rows)
    
    _add_heading(doc, "4-2 العضوية في المنظمات", size=13)
    df_intl = load_q_list("Q_عضويات_دولية")
    df_local = load_q_list("Q_عضويات_محلية")
    max_len = max(len(df_intl), len(df_local)) if (not df_intl.empty or not df_local.empty) else 0
    rows = []
    for i in range(max_len):
        intl_val = df_intl.iloc[i]["المنظمة"] if not df_intl.empty and i < len(df_intl) else ""
        local_val = df_local.iloc[i]["المنظمة"] if not df_local.empty and i < len(df_local) else ""
        rows.append([intl_val, local_val])
    _add_table(doc, ["عضوية في منظمات دولية", "عضوية في منظمات محلية"], rows)
    doc.add_page_break()
    
    # ============ 5. التدريب ============
    _add_heading(doc, "5. التدريب / التعليم (8%)", size=16)
    _add_para(doc, "يتم التركيز على برامج التدريب وورش العمل التي تمت خلال الفترة.")
    
    rows = [
        ["دورات بالداخل", str(qval("دورات بالداخل عدد")), str(qval("دورات بالداخل متدربين"))],
        ["دورات بالخارج", str(qval("دورات بالخارج عدد")), str(qval("دورات بالخارج متدربين"))],
        ["دورات online", str(qval("دورات online عدد")), str(qval("دورات online متدربين"))],
        ["ورش عمل/ندوات بالداخل", str(qval("ورش داخل عدد")), str(qval("ورش داخل متدربين"))],
        ["ورش عمل/ندوات بالخارج", str(qval("ورش خارج عدد")), str(qval("ورش خارج متدربين"))],
        ["مؤتمرات بالداخل", str(qval("مؤتمرات داخل عدد")), str(qval("مؤتمرات داخل متدربين"))],
        ["مؤتمرات بالخارج", str(qval("مؤتمرات خارج عدد")), str(qval("مؤتمرات خارج متدربين"))],
        ["دورات نظمها المعهد", str(qval("دورات نظمها المعهد عدد")), str(qval("دورات نظمها المعهد متدربين"))],
    ]
    total_cnt = sum(int(r[1]) for r in rows)
    total_trn = sum(int(r[2]) for r in rows)
    rows.append(["الإجمالي", str(total_cnt), str(total_trn)])
    _add_table(doc, ["الدورة/ورشة العمل", "عدد الدورات", "عدد المتدربين"], rows)
    doc.add_page_break()
    
    # ============ 6. التمويل ============
    _add_heading(doc, "6. التمويل والأداء المالي (15%)", size=16)
    rows = [
        ["تمويل الخطة البحثية من إجمالي التمويل", str(qvalf("تمويل الخطة مبلغ")), f"{qvalf('تمويل الخطة نسبة'):.1f}%"],
        ["التمويل الخارجي من إجمالي التمويل", str(qvalf("تمويل خارجي مبلغ")), f"{qvalf('تمويل خارجي نسبة'):.1f}%"],
        ["المنصرف الفعلي", str(qvalf("منصرف فعلي مبلغ")), f"{qvalf('منصرف فعلي نسبة'):.1f}%"],
        ["تمويل النشر العلمي - دراسات (ماجستير/دكتوراه)", str(qvalf("تمويل نشر دراسات مبلغ")), "-"],
        ["تمويل النشر العلمي - أبحاث", str(qvalf("تمويل نشر أبحاث مبلغ")), "-"],
        ["تمويل النشر العلمي - مكافآت", str(qvalf("تمويل نشر مكافآت مبلغ")), "-"],
    ]
    _add_table(doc, ["البند", "المبلغ (بألاف جنيه)", "النسبة"], rows)
    doc.add_page_break()
    
    # ============ 7. الموارد البشرية ============
    _add_heading(doc, "7. مؤشرات الأداء للموارد البشرية (10%)", size=16)
    rows = [
        ["عدد الاجتماعات الدورية مع الهيئة البحثية والعاملين", str(qval("عدد الاجتماعات الدورية"))],
        ["عدد اجتماعات مجلس إدارة المعهد", str(qval("عدد اجتماعات مجلس الإدارة"))],
        ["عدد اجتماعات وحدة الفكر بالمعهد", str(qval("عدد اجتماعات وحدة الفكر"))],
        ["الآليات المستخدمة لتسويق إمكانات المعهد", qstr("آليات التسويق")],
        ["آليات الثواب والعقاب لتحفيز الباحثين والعاملين", qstr("آليات الثواب والعقاب")],
    ]
    _add_table(doc, ["البند", "القيمة"], rows)
    doc.add_page_break()
    
    # ============ 8. الاجتماعات الوزارية ============
    _add_heading(doc, "8. الاجتماعات الوزارية", size=16)
    df = load_q_list("Q_اجتماعات_وزارية")
    rows = [[r.get("التاريخ",""), r.get("الاسم",""), r.get("الوظيفة",""),
            r.get("مكان انعقاد الاجتماع",""), r.get("الغرض من الاجتماع","")] for _, r in df.iterrows()] if not df.empty else []
    _add_table(doc, ["التاريخ", "الاسم", "الوظيفة", "مكان انعقاد الاجتماع", "الغرض من الاجتماع"], rows)
    
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()

# ============================================================
# التقرير السنوي (Excel)
# ============================================================
def generate_annual_report_excel(institute, fiscal_year):
    wb = Workbook()
    ws = wb.active
    ws.title = f"تقرير سنوي {fiscal_year}".replace("/", "_")
    ws.sheet_view.rightToLeft = True
    
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    data_font = Font(name='Arial', size=10)
    data_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    fy_s = int(fiscal_year.split("/")[0]); fy_e = int(fiscal_year.split("/")[1])
    kpi_df = load_sheet("KPI_Data", KPI_COLUMNS)
    year_df = pd.DataFrame()
    if not kpi_df.empty:
        kpi_f = kpi_df[kpi_df["المعهد"] == institute]
        fy_months = [("يوليو", fy_s), ("أغسطس", fy_s), ("سبتمبر", fy_s), ("أكتوبر", fy_s), ("نوفمبر", fy_s), ("ديسمبر", fy_s),
                     ("يناير", fy_e), ("فبراير", fy_e), ("مارس", fy_e), ("أبريل", fy_e), ("مايو", fy_e), ("يونيو", fy_e)]
        mask = pd.Series([False]*len(kpi_f), index=kpi_f.index)
        for m, y in fy_months:
            mask |= ((kpi_f["الشهر"] == m) & (kpi_f["السنة"] == str(y)))
        year_df = kpi_f[mask].copy()
        for c in year_df.columns:
            if c not in ["المعهد","المستخدم","الشهر","السنة"]:
                year_df[c] = pd.to_numeric(year_df[c], errors="coerce").fillna(0)
    agg = year_df.sum(numeric_only=True) if not year_df.empty else pd.Series()
    
    ws.merge_cells('A1:B1')
    ws['A1'] = f"تقرير مؤشرات الأداء السنوي - {institute}"
    ws['A1'].font = Font(name='Arial', size=16, bold=True)
    ws['A1'].alignment = center_align
    ws.merge_cells('A2:B2')
    ws['A2'] = f"السنة المالية {fiscal_year}"
    ws['A2'].font = Font(name='Arial', size=13, bold=True)
    ws['A2'].alignment = center_align
    
    row = 4
    for col, lbl in [(1, "المؤشر"), (2, "القيمة")]:
        ws.cell(row=row, column=col, value=lbl).font = header_font
        ws.cell(row=row, column=col).fill = header_fill
        ws.cell(row=row, column=col).alignment = center_align
        ws.cell(row=row, column=col).border = border
    row += 1
    
    items = [
        ("إجمالي الدراسات الجارية", int(agg.get("إجمالي الدراسات الجارية", 0)) if not agg.empty else 0),
        ("خطة بحثية", int(agg.get("خطة بحثية", 0)) if not agg.empty else 0),
        ("استشارية", int(agg.get("استشارية", 0)) if not agg.empty else 0),
        ("دراسات بتمويل ذاتي", int(agg.get("دراسات بتمويل ذاتي", 0)) if not agg.empty else 0),
        ("إجمالي التقارير", int(agg.get("إجمالي التقارير", 0)) if not agg.empty else 0),
        ("عدد المذكرات", int(agg.get("عدد المذكرات", 0)) if not agg.empty else 0),
        ("إجمالي المقترحات", int(agg.get("إجمالي المقترحات", 0)) if not agg.empty else 0),
        ("إجمالي المأموريات", int(agg.get("إجمالي عدد المأموريات", 0)) if not agg.empty else 0),
        ("إجمالي أيام المأموريات", int(agg.get("إجمالي عدد الأيام الفعلية", 0)) if not agg.empty else 0),
        ("إجمالي المعاينات", int(agg.get("إجمالي عدد المعاينات", 0)) if not agg.empty else 0),
        ("متدربين", int(agg.get("متدربين", 0)) if not agg.empty else 0),
        ("إجمالي المشاركين في أنشطة", int(agg.get("إجمالي المشاركين في أنشطة علمية", 0)) if not agg.empty else 0),
        ("إجمالي الاجتماعات", int(agg.get("إجمالي الاجتماعات", 0)) if not agg.empty else 0),
        ("إجمالي الأوراق", int(agg.get("إجمالي الأوراق البحثية", 0)) if not agg.empty else 0),
        ("إجمالي الرسائل", int(agg.get("إجمالي الرسائل", 0)) if not agg.empty else 0),
    ]
    for label, val in items:
        ws.cell(row=row, column=1, value=label).font = data_font
        ws.cell(row=row, column=1).alignment = data_align
        ws.cell(row=row, column=1).border = border
        ws.cell(row=row, column=2, value=val).font = data_font
        ws.cell(row=row, column=2).alignment = center_align
        ws.cell(row=row, column=2).border = border
        row += 1
    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 20
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# ============================================================
# المستخدمون
# ============================================================
USERS = {
    "admin": {"password": "admin123", "role": "admin", "institute": "الجميع"},
    "wmri": {"password": "1234", "role": "user", "institute": "معهد بحوث إدارة المياه"},
    "dri": {"password": "1234", "role": "user", "institute": "معهد بحوث الصرف"},
    "wrri": {"password": "1234", "role": "user", "institute": "معهد بحوث الموارد المائية"},
    "nri": {"password": "1234", "role": "user", "institute": "معهد بحوث النيل"},
    "hri": {"password": "1234", "role": "user", "institute": "معهد بحوث الهيدروليكا"},
    "cori": {"password": "1234", "role": "user", "institute": "معهد بحوث الشواطئ"},
    "gwri": {"password": "1234", "role": "user", "institute": "معهد بحوث المياه الجوفية"},
    "chri": {"password": "1234", "role": "user", "institute": "معهد بحوث صيانة القنوات"},
    "eri": {"password": "1234", "role": "user", "institute": "معهد بحوث الإنشاءات"},
    "mri": {"password": "1234", "role": "user", "institute": "معهد بحوث الميكانيكا"},
    "sri": {"password": "1234", "role": "user", "institute": "معهد بحوث المساحة"},
    "ecri": {"password": "1234", "role": "user", "institute": "معهد البيئة وتغير المناخ"},
}

for k, v in [("logged_in", False), ("gs_bust", 0), ("data_bust", 0), ("editor_counter", 0), ("q_editor_counter", 0)]:
    if k not in st.session_state:
        st.session_state[k] = v

# ============================================================
# تسجيل الدخول
# ============================================================
if not st.session_state.logged_in:
    st.title("🔐 تسجيل الدخول - المركز القومي لبحوث المياه")
    c1, c2, c3 = st.columns([1, 1, 1])
    with c2:
        username = st.text_input("اسم المستخدم")
        password = st.text_input("كلمة المرور", type="password")
        if st.button("🚪 دخول", type="primary"):
            if username in USERS and USERS[username]["password"] == password:
                st.session_state.logged_in = True
                st.session_state.username = username
                st.session_state.role = USERS[username]["role"]
                st.session_state.institute = USERS[username]["institute"]
                reset_all()
                st.rerun()
            else:
                st.error("❌ اسم المستخدم أو كلمة المرور غير صحيحة")
    st.stop()

# ============================================================
# بعد الدخول
# ============================================================
c_head = st.columns([5, 1, 1])
with c_head[0]:
    st.success(f"👤 مرحباً **{st.session_state.username}** — 🏢 {st.session_state.institute}")
with c_head[1]:
    if st.button("🔄 تحديث"):
        reset_all(); st.rerun()
with c_head[2]:
    if st.button("🚪 خروج"):
        st.session_state.clear()
        st.cache_data.clear(); st.cache_resource.clear()
        st.rerun()

# ============================================================
# ⭐ التبويبات الجديدة
# ============================================================
if st.session_state.role == "admin":
    tab_m_dash, tab_q_dash, tab_annual, tab_data, tab_debug = st.tabs([
        "📊 Dashboard الشهري", "📈 التقرير الربع سنوي", "📅 التقرير السنوي", "📄 البيانات", "🔧 تشخيص"])
    tab_m_input = None
    tab_q_input = None
else:
    tab_m_input, tab_q_input, tab_m_dash, tab_q_dash, tab_annual, tab_data = st.tabs([
        "📥 إدخال البيانات الشهرية",
        "📊 إدخال البيانات الربع سنوية",
        "📊 Dashboard الشهري",
        "📈 التقرير الربع سنوي",
        "📅 التقرير السنوي",
        "📄 البيانات",
    ])
    tab_debug = None

# ============================================================
# 📥 تبويب الإدخال الشهري (مع النصف سنوي والسنوي)
# ============================================================
if tab_m_input is not None:
    with tab_m_input:
        st.header("📥 إدخال البيانات الشهرية")
        
        c1, c2 = st.columns(2)
        with c1:
            month = st.selectbox("الشهر", MONTHS, key="m_month")
        with c2:
            year = st.number_input("السنة", min_value=2020, max_value=2035, value=2026, key="m_year")
        
        fp = get_fiscal_period(month, year)
        period_val = f"{month} {year}"
        
        info_cols = st.columns(4)
        info_cols[0].info(f"🗓️ **{month} {year}**")
        info_cols[1].info(f"💰 السنة المالية: **{fp['fiscal_year']}**")
        info_cols[2].info(f"📊 النصف: **{fp['half'] if fp['half'] else 'لا'}")
        info_cols[3].info(f"📅 سنوي: **{'نعم' if fp['is_annual_month'] else 'لا'}**")
        
        st.divider()
        st.markdown("## 📋 القوائم الشهرية")
        st.caption("يتم احتساب الأرقام تلقائياً. الأعمدة ذات القوائم المنسدلة محددة بـ ▾")
        
        list_tabs = st.tabs([cfg["title"] for cfg in MONTHLY_LISTS.values()])
        
        for idx, (sheet_name, cfg) in enumerate(MONTHLY_LISTS.items()):
            with list_tabs[idx]:
                st.markdown(f"### {cfg['title']}")
                if cfg["hint"]:
                    st.info(f"💡 {cfg['hint']}")
                
                save_key = f"_saved_{sheet_name}_{st.session_state.institute}_{period_val}"
                if save_key in st.session_state:
                    data_clean = st.session_state[save_key].copy()
                    for col in cfg["cols"]:
                        if col not in data_clean.columns:
                            data_clean[col] = ""
                    data_clean = data_clean[cfg["cols"]]
                    st.info(f"💾 {len(data_clean)} صف من الذاكرة المؤقتة")
                else:
                    full_cols = cfg["cols"] + ["المعهد", "الفترة"]
                    data = load_sheet(sheet_name, full_cols)
                    if not data.empty and "المعهد" in data.columns and "الفترة" in data.columns:
                        data = data[(data["المعهد"] == st.session_state.institute) & (data["الفترة"] == period_val)]
                    if data.empty:
                        data_clean = pd.DataFrame(columns=cfg["cols"])
                        st.caption("📝 لا توجد بيانات محفوظة")
                    else:
                        data_clean = data.reindex(columns=cfg["cols"]).fillna("").reset_index(drop=True)
                        st.success(f"✅ تم استرجاع {len(data_clean)} صف")
                
                editor_key = f"{sheet_name}_ed_{period_val}_{st.session_state.editor_counter}"
                display_df = rtl_display(data_clean) if not data_clean.empty else pd.DataFrame(columns=list(reversed(cfg["cols"])))
                
                column_config = {}
                for col_name, options in cfg.get("dropdowns", {}).items():
                    column_config[col_name] = st.column_config.SelectboxColumn(
                        col_name, options=options, required=False, help=f"اختر من: {' / '.join(options)}")
                
                input_rtl = st.data_editor(display_df, num_rows="dynamic", use_container_width=True,
                                            key=editor_key, hide_index=True, column_config=column_config)
                
                if not input_rtl.empty:
                    input_df = input_rtl.reindex(columns=cfg["cols"]).fillna("")
                else:
                    input_df = pd.DataFrame(columns=cfg["cols"])
                
                if st.button(f"💾 حفظ {cfg['title']}", key=f"{sheet_name}_sv"):
                    if save_list(sheet_name, input_df, cfg["cols"], st.session_state.institute, "الفترة", period_val):
                        st.session_state.editor_counter = st.session_state.get("editor_counter", 0) + 1
                        st.success(f"✅ تم الحفظ")
                        st.balloons()
                        time.sleep(0.3)
                        st.rerun()
        
        # ⭐ النصف سنوي (يناير، يوليو)
        if fp["is_half_month"]:
            st.divider()
            st.markdown(f"## 📊 المدخلات النصف سنوية — {fp['half']}")
            st.caption("تظهر هذه المدخلات في شهري يناير ويوليو فقط")
            
            df_h = load_sheet("البيانات_النصف_سنوية", HALF_COLUMNS)
            df_h_f = pd.DataFrame()
            if not df_h.empty and "المعهد" in df_h.columns:
                df_h_f = df_h[(df_h["المعهد"] == st.session_state.institute) & (df_h["النصف"] == fp["half"]) & (df_h["السنة المالية"] == fp["fiscal_year"])]
            existing_h = df_h_f.iloc[0] if not df_h_f.empty else None
            
            def hv(col, default=0):
                if existing_h is None or col not in existing_h.index: return default
                try: return int(float(existing_h[col])) if existing_h[col] and str(existing_h[col]) not in ["", "nan"] else default
                except: return default
            
            def hvf(col, default=0.0):
                if existing_h is None or col not in existing_h.index: return default
                try: return float(existing_h[col]) if existing_h[col] and str(existing_h[col]) not in ["", "nan"] else default
                except: return default
            
            kp_h = f"{fp['half']}_{fp['fiscal_year']}".replace("/", "_").replace(" ", "_")
            c1, c2 = st.columns(2)
            h_studies = c1.number_input("إجمالي دراسات الخطة المنجزة", min_value=0, value=hv("إجمالي عدد دراسات الخطة البحثية المنجزة"), key=f"h_st_{kp_h}")
            h_outputs = c2.number_input("إجمالي المخرجات التطبيقية", min_value=0, value=hv("إجمالي عدد المخرجات التطبيقية من دراسات ومشروعات"), key=f"h_out_{kp_h}")
            h_revenue = st.number_input("إجمالي الإيرادات الذاتية", min_value=0.0, value=hvf("إجمالي مبلغ الإيرادات الذاتية"), format="%.2f", key=f"h_rev_{kp_h}")
            
            if st.button("💾 حفظ البيانات النصف سنوية", type="primary", key="save_h"):
                h_data = {"المعهد": st.session_state.institute, "المستخدم": st.session_state.username,
                          "النصف": fp["half"], "السنة المالية": fp["fiscal_year"],
                          "إجمالي عدد دراسات الخطة البحثية المنجزة": h_studies,
                          "إجمالي عدد المخرجات التطبيقية من دراسات ومشروعات": h_outputs,
                          "إجمالي مبلغ الإيرادات الذاتية": h_revenue}
                result = save_or_update_row("البيانات_النصف_سنوية", HALF_COLUMNS, h_data, ["المعهد", "النصف", "السنة المالية"])
                if result:
                    invalidate_all()
                    st.success(f"✅ تم {'تحديث' if result == 'updated' else 'حفظ'} البيانات النصف سنوية")
                    time.sleep(0.5); st.rerun()
        
        # ⭐ السنوي (يوليو فقط)
        if fp["is_annual_month"]:
            st.divider()
            st.markdown("## 📊 المدخلات السنوية")
            st.caption("تظهر هذه المدخلات في شهر يوليو فقط")
            
            df_a = load_sheet("البيانات_السنوية", ANNUAL_COLUMNS)
            df_a_f = pd.DataFrame()
            if not df_a.empty and "المعهد" in df_a.columns:
                df_a_f = df_a[(df_a["المعهد"] == st.session_state.institute) & (df_a["السنة المالية"] == fp["fiscal_year"])]
            existing_a = df_a_f.iloc[0] if not df_a_f.empty else None
            
            def av(col, default=0):
                if existing_a is None or col not in existing_a.index: return default
                try: return int(float(existing_a[col])) if existing_a[col] and str(existing_a[col]) not in ["", "nan"] else default
                except: return default
            
            a_patents = st.number_input("عدد براءات الاختراع", min_value=0, value=av("عدد براءات الاختراع"), key=f"a_pat_{fp['fiscal_year'].replace('/', '_')}")
            if st.button("💾 حفظ البيانات السنوية", type="primary", key="save_a"):
                a_data = {"المعهد": st.session_state.institute, "المستخدم": st.session_state.username,
                          "السنة المالية": fp["fiscal_year"], "عدد براءات الاختراع": a_patents}
                result = save_or_update_row("البيانات_السنوية", ANNUAL_COLUMNS, a_data, ["المعهد", "السنة المالية"])
                if result:
                    invalidate_all()
                    st.success(f"✅ تم {'تحديث' if result == 'updated' else 'حفظ'} البيانات السنوية")
                    time.sleep(0.5); st.rerun()
        
        # المعاينة
        st.divider()
        st.markdown("## 📊 معاينة الأرقام المحسوبة تلقائياً")
        kpi_preview = compute_kpi_from_lists(st.session_state.institute, month, year)
        
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("📚 الدراسات", kpi_preview["إجمالي الدراسات الجارية"])
        c2.metric("📄 التقارير", kpi_preview["إجمالي التقارير"])
        c3.metric("📝 المذكرات", kpi_preview["عدد المذكرات"])
        c4.metric("📋 المقترحات", kpi_preview["إجمالي المقترحات"])
        
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("🏃 المأموريات", kpi_preview["إجمالي عدد المأموريات"])
        c2.metric("🔍 المعاينات", kpi_preview["إجمالي عدد المعاينات"])
        c3.metric("👨‍🎓 المتدربين", kpi_preview["متدربين"])
        c4.metric("🤝 الاجتماعات", kpi_preview["إجمالي الاجتماعات"])
        
        with st.expander("📋 تفاصيل جميع المؤشرات"):
            detail_df = pd.DataFrame([kpi_preview]).T
            detail_df.columns = ["القيمة"]; detail_df.index.name = "المؤشر"
            st.dataframe(detail_df, use_container_width=True)
        
        if st.button("💾 حفظ/تحديث مؤشرات الأداء الشهرية", type="primary", key="save_kpi"):
            kpi_data = {"المعهد": st.session_state.institute, "المستخدم": st.session_state.username,
                        "الشهر": month, "السنة": int(year)}
            kpi_data.update(kpi_preview)
            result = save_or_update_row("KPI_Data", KPI_COLUMNS, kpi_data, ["المعهد", "الشهر", "السنة"])
            if result:
                invalidate_all()
                st.success(f"✅ تم {'تحديث' if result == 'updated' else 'حفظ'} مؤشرات الأداء")
                st.balloons()
                time.sleep(0.5); st.rerun()
        
        st.divider()
        st.markdown("## 📥 تصدير التقرير الشهري")
        if st.button("📄 توليد التقرير الشهري (Word)", key="gen_monthly_docx"):
            with st.spinner("جاري التوليد..."):
                try:
                    docx_bytes = generate_monthly_report(st.session_state.institute, month, year)
                    st.download_button(
                        label="⬇️ تحميل التقرير الشهري",
                        data=docx_bytes,
                        file_name=f"تقرير_{st.session_state.institute}_{month}_{year}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="dl_monthly")
                except Exception as e:
                    st.error(f"خطأ: {e}")

# ============================================================
# 📊 تبويب الإدخال الربع سنوي (مع توليد التقرير)
# ============================================================
if tab_q_input is not None:
    with tab_q_input:
        st.header("📊 إدخال البيانات الربع سنوية")
        
        c1, c2 = st.columns(2)
        with c1:
            quarter = st.selectbox("الربع", ["الربع الأول", "الربع الثاني", "الربع الثالث", "الربع الرابع"], key="q_quarter")
        with c2:
            fiscal_year = st.selectbox("السنة المالية", ["2025/2026", "2026/2027", "2027/2028"], key="q_fy")
        
        q_months = get_quarter_months(quarter, fiscal_year)
        st.info(f"📅 يشمل التقرير: **{' + '.join([f'{m} {y}' for m, y in q_months])}**")
        st.divider()
        
        period_q = f"{quarter} - {fiscal_year}"
        
        # تحميل البيانات الربع سنوية الحالية
        df_qk = load_sheet("Quarterly_KPI", QUARTERLY_KPI_COLUMNS)
        q_existing = None
        if not df_qk.empty and "الربع" in df_qk.columns:
            r = df_qk[(df_qk["المعهد"] == st.session_state.institute) & (df_qk["الربع"] == quarter) & (df_qk["السنة المالية"] == fiscal_year)]
            if not r.empty:
                q_existing = r.iloc[0]
        
        def qv(col, default=0):
            if q_existing is None or col not in q_existing.index: return default
            try: return int(float(q_existing[col])) if q_existing[col] and str(q_existing[col]) not in ["", "nan"] else default
            except: return default
        
        def qvf(col, default=0.0):
            if q_existing is None or col not in q_existing.index: return default
            try: return float(q_existing[col]) if q_existing[col] and str(q_existing[col]) not in ["", "nan"] else default
            except: return default
        
        def qs(col, default=""):
            if q_existing is None or col not in q_existing.index: return default
            return str(q_existing[col]) if pd.notna(q_existing[col]) else default
        
        if q_existing is not None:
            st.warning(f"ℹ️ توجد بيانات محفوظة لهذا الربع - سيتم تحديثها عند الحفظ")
        
        # تحضير البيانات المحفوظة مؤقتاً
        q_editor_key_suffix = f"{quarter}_{fiscal_year}".replace("/", "_").replace(" ", "_")
        
        # التبويبات الفرعية للقوائم التفصيلية
        st.markdown("## 📋 القوائم التفصيلية للبيانات الربع سنوية")
        
        q_tabs = st.tabs([
            "1-1-1 فردية", "1-1-2 مشتركة", "2-1-1 ممولة محلياً", "2-1-2 ممولة دولياً",
            "2-1-3 تكليفات", "2-1-4 استشارات وزارة", "2-1-5 استشارات جهات أخرى",
            "2-2-1 مقترحات محلية", "2-2-2 مقترحات دولية", "2-2-3 مقترحات للوزارة", "2-2-4 مقترحات خارج الوزارة",
            "3-1 تعاون معاهد", "3-2 تعاون جهات", "3-3 تعاون خارجي",
            "4-2 عضويات دولية", "4-2 عضويات محلية", "8 اجتماعات وزارية"
        ])
        
        for idx, (sheet_name, cfg) in enumerate(QUARTERLY_LISTS.items()):
            with q_tabs[idx]:
                st.markdown(f"### {cfg['title']}")
                
                save_key = f"_saved_{sheet_name}_{st.session_state.institute}_{period_q}"
                if save_key in st.session_state:
                    data_clean = st.session_state[save_key].copy()
                    for col in cfg["cols"]:
                        if col not in data_clean.columns:
                            data_clean[col] = ""
                    data_clean = data_clean[cfg["cols"]]
                    st.info(f"💾 {len(data_clean)} صف من الذاكرة المؤقتة")
                else:
                    full_cols = cfg["cols"] + ["المعهد", "الفترة"]
                    data = load_sheet(sheet_name, full_cols)
                    if not data.empty and "المعهد" in data.columns and "الفترة" in data.columns:
                        data = data[(data["المعهد"] == st.session_state.institute) & (data["الفترة"] == period_q)]
                    if data.empty:
                        data_clean = pd.DataFrame(columns=cfg["cols"])
                        st.caption("📝 لا توجد بيانات محفوظة")
                    else:
                        data_clean = data.reindex(columns=cfg["cols"]).fillna("").reset_index(drop=True)
                        st.success(f"✅ تم استرجاع {len(data_clean)} صف")
                
                editor_key = f"q_{sheet_name}_ed_{q_editor_key_suffix}_{st.session_state.get('q_editor_counter', 0)}"
                display_df = rtl_display(data_clean) if not data_clean.empty else pd.DataFrame(columns=list(reversed(cfg["cols"])))
                
                input_rtl = st.data_editor(display_df, num_rows="dynamic", use_container_width=True,
                                            key=editor_key, hide_index=True)
                
                if not input_rtl.empty:
                    input_df = input_rtl.reindex(columns=cfg["cols"]).fillna("")
                else:
                    input_df = pd.DataFrame(columns=cfg["cols"])
                
                if st.button(f"💾 حفظ {cfg['title']}", key=f"q_{sheet_name}_sv"):
                    if save_list(sheet_name, input_df, cfg["cols"], st.session_state.institute, "الفترة", period_q):
                        st.session_state.q_editor_counter = st.session_state.get("q_editor_counter", 0) + 1
                        st.success(f"✅ تم الحفظ")
                        st.balloons()
                        time.sleep(0.3)
                        st.rerun()
        
        # الأرقام الربع سنوية
        st.divider()
        st.markdown("## 📊 الأرقام الربع سنوية")
        
        with st.expander("1. الأداء التنفيذي للخطة البحثية", expanded=False):
            c1, c2, c3 = st.columns(3)
            q_ind_count = c1.number_input("عدد المشاريع الفردية", min_value=0, value=qv("عدد المشاريع الفردية"), key="q_ind_c")
            q_ind_pct = c2.number_input("نسبة تنفيذ المشاريع الفردية %", min_value=0.0, max_value=100.0, value=qvf("نسبة تنفيذ المشاريع الفردية"), format="%.1f", key="q_ind_p")
            
            c1, c2, c3 = st.columns(3)
            q_jnt_count = c1.number_input("عدد المشاريع المشتركة", min_value=0, value=qv("عدد المشاريع المشتركة"), key="q_jnt_c")
            q_jnt_pct = c2.number_input("نسبة تنفيذ المشاريع المشتركة %", min_value=0.0, max_value=100.0, value=qvf("نسبة تنفيذ المشاريع المشتركة"), format="%.1f", key="q_jnt_p")
            q_jnt_share = c3.number_input("نسبة المشاركة %", min_value=0.0, max_value=100.0, value=qvf("نسبة المشاركة المشتركة"), format="%.1f", key="q_jnt_s")
        
        with st.expander("2. الاستشارات والدراسات خارج الخطة"):
            c1, c2 = st.columns(2)
            q_local_proj = c1.number_input("مشاريع ممولة محلياً", min_value=0, value=qv("عدد المشاريع الممولة محلياً"), key="q_lp")
            q_intl_proj = c2.number_input("مشاريع ممولة دولياً", min_value=0, value=qv("عدد المشاريع الممولة دولياً"), key="q_ip")
            
            c1, c2, c3 = st.columns(3)
            q_ministry_assign = c1.number_input("تكليفات من الوزارة", min_value=0, value=qv("عدد التكليفات من الوزارة"), key="q_ma")
            q_ministry_consult = c2.number_input("استشارات من الوزارة", min_value=0, value=qv("عدد الاستشارات من الوزارة"), key="q_mc")
            q_other_consult = c3.number_input("استشارات جهات حكومية/خاصة", min_value=0, value=qv("عدد الاستشارات من جهات حكومية/خاصة"), key="q_oc")
            
            c1, c2, c3, c4 = st.columns(4)
            q_prop_local = c1.number_input("مقترحات محلية", min_value=0, value=qv("عدد المقترحات محلياً"), key="q_pl")
            q_prop_intl = c2.number_input("مقترحات دولية", min_value=0, value=qv("عدد المقترحات دولياً"), key="q_pi")
            q_prop_ministry = c3.number_input("مقترحات للوزارة", min_value=0, value=qv("عدد المقترحات للوزارة"), key="q_pm")
            q_prop_outside = c4.number_input("مقترحات خارج الوزارة", min_value=0, value=qv("عدد المقترحات خارج الوزارة"), key="q_po")
        
        with st.expander("3. التعاون"):
            c1, c2, c3 = st.columns(3)
            q_coop_inst = c1.number_input("تعاون مع المعاهد", min_value=0, value=qv("عدد تعاون مع المعاهد"), key="q_ci")
            q_coop_entities = c2.number_input("تعاون مع جهات خارج المعاهد", min_value=0, value=qv("عدد تعاون مع جهات خارج المعاهد"), key="q_ce")
            q_coop_foreign = c3.number_input("تعاون خارجي", min_value=0, value=qv("عدد تعاون خارجي"), key="q_cf")
        
        with st.expander("4. الإنتاج العلمي والنشر"):
            c1, c2, c3, c4 = st.columns(4)
            q_pub = c1.number_input("أبحاث منشورة", min_value=0, value=qv("عدد الأبحاث المنشورة"), key="q_pub")
            q_q1 = c2.number_input("أبحاث Q1/Q2", min_value=0, value=qv("عدد الأبحاث Q1/Q2"), key="q_q1")
            q_appr = c3.number_input("أبحاث بموافقة نشر", min_value=0, value=qv("عدد الأبحاث بموافقة نشر"), key="q_appr")
            q_rew = c4.number_input("أبحاث بمكافأة نشر", min_value=0, value=qv("عدد الأبحاث بمكافأة نشر"), key="q_rew")
            
            c1, c2, c3 = st.columns(3)
            q_res_rew = c1.number_input("باحثون بمكافأة", min_value=0, value=qv("عدد الباحثين بمكافأة"), key="q_rrw")
            q_mas = c2.number_input("رسائل ماجستير", min_value=0, value=qv("عدد رسائل الماجستير"), key="q_mas")
            q_phd = c3.number_input("رسائل دكتوراه", min_value=0, value=qv("عدد رسائل الدكتوراه"), key="q_phd")
            
            c1, c2, c3 = st.columns(3)
            q_review = c1.number_input("مراجعة بمجلات", min_value=0, value=qv("مراجعة وتحكيم بمجلات"), key="q_rev")
            q_sup_in = c2.number_input("إشراف لباحثي المعهد", min_value=0, value=qv("إشراف لباحثي المعهد"), key="q_si")
            q_sup_out = c3.number_input("إشراف لباحثين خارجيين", min_value=0, value=qv("إشراف لباحثين خارجيين"), key="q_so")
            
            c1, c2 = st.columns(2)
            q_mem_intl = c1.number_input("عضويات دولية", min_value=0, value=qv("عدد العضويات الدولية"), key="q_mi")
            q_mem_local = c2.number_input("عضويات محلية", min_value=0, value=qv("عدد العضويات المحلية"), key="q_ml")
        
        with st.expander("5. التدريب"):
            c1, c2 = st.columns(2)
            q_tr_in_c = c1.number_input("دورات بالداخل - عدد", min_value=0, value=qv("دورات بالداخل عدد"), key="q_tic")
            q_tr_in_t = c2.number_input("دورات بالداخل - متدربين", min_value=0, value=qv("دورات بالداخل متدربين"), key="q_tit")
            
            c1, c2 = st.columns(2)
            q_tr_out_c = c1.number_input("دورات بالخارج - عدد", min_value=0, value=qv("دورات بالخارج عدد"), key="q_toc")
            q_tr_out_t = c2.number_input("دورات بالخارج - متدربين", min_value=0, value=qv("دورات بالخارج متدربين"), key="q_tot")
            
            c1, c2 = st.columns(2)
            q_tr_onl_c = c1.number_input("دورات online - عدد", min_value=0, value=qv("دورات online عدد"), key="q_tonc")
            q_tr_onl_t = c2.number_input("دورات online - متدربين", min_value=0, value=qv("دورات online متدربين"), key="q_tont")
            
            c1, c2 = st.columns(2)
            q_ws_in_c = c1.number_input("ورش داخل - عدد", min_value=0, value=qv("ورش داخل عدد"), key="q_wic")
            q_ws_in_t = c2.number_input("ورش داخل - متدربين", min_value=0, value=qv("ورش داخل متدربين"), key="q_wit")
            
            c1, c2 = st.columns(2)
            q_ws_out_c = c1.number_input("ورش خارج - عدد", min_value=0, value=qv("ورش خارج عدد"), key="q_woc")
            q_ws_out_t = c2.number_input("ورش خارج - متدربين", min_value=0, value=qv("ورش خارج متدربين"), key="q_wot")
            
            c1, c2 = st.columns(2)
            q_cf_in_c = c1.number_input("مؤتمرات داخل - عدد", min_value=0, value=qv("مؤتمرات داخل عدد"), key="q_cic")
            q_cf_in_t = c2.number_input("مؤتمرات داخل - متدربين", min_value=0, value=qv("مؤتمرات داخل متدربين"), key="q_cit")
            
            c1, c2 = st.columns(2)
            q_cf_out_c = c1.number_input("مؤتمرات خارج - عدد", min_value=0, value=qv("مؤتمرات خارج عدد"), key="q_coc")
            q_cf_out_t = c2.number_input("مؤتمرات خارج - متدربين", min_value=0, value=qv("مؤتمرات خارج متدربين"), key="q_cot")
            
            c1, c2 = st.columns(2)
            q_org_c = c1.number_input("دورات نظمها المعهد - عدد", min_value=0, value=qv("دورات نظمها المعهد عدد"), key="q_orgc")
            q_org_t = c2.number_input("دورات نظمها المعهد - متدربين", min_value=0, value=qv("دورات نظمها المعهد متدربين"), key="q_orgt")
        
        with st.expander("6. التمويل"):
            c1, c2 = st.columns(2)
            q_fund_plan_amt = c1.number_input("تمويل الخطة مبلغ", min_value=0.0, value=qvf("تمويل الخطة مبلغ"), format="%.2f", key="q_fpa")
            q_fund_plan_pct = c2.number_input("تمويل الخطة نسبة %", min_value=0.0, max_value=100.0, value=qvf("تمويل الخطة نسبة"), format="%.1f", key="q_fpp")
            
            c1, c2 = st.columns(2)
            q_fund_ext_amt = c1.number_input("التمويل الخارجي مبلغ", min_value=0.0, value=qvf("تمويل خارجي مبلغ"), format="%.2f", key="q_fea")
            q_fund_ext_pct = c2.number_input("التمويل الخارجي نسبة %", min_value=0.0, max_value=100.0, value=qvf("تمويل خارجي نسبة"), format="%.1f", key="q_fep")
            
            c1, c2 = st.columns(2)
            q_spend_amt = c1.number_input("المنصرف الفعلي مبلغ", min_value=0.0, value=qvf("منصرف فعلي مبلغ"), format="%.2f", key="q_spa")
            q_spend_pct = c2.number_input("المنصرف الفعلي نسبة %", min_value=0.0, max_value=200.0, value=qvf("منصرف فعلي نسبة"), format="%.1f", key="q_spp")
            
            c1, c2, c3 = st.columns(3)
            q_pub_th = c1.number_input("تمويل النشر - دراسات", min_value=0.0, value=qvf("تمويل نشر دراسات مبلغ"), format="%.2f", key="q_pth")
            q_pub_rs = c2.number_input("تمويل النشر - أبحاث", min_value=0.0, value=qvf("تمويل نشر أبحاث مبلغ"), format="%.2f", key="q_prs")
            q_pub_rw = c3.number_input("تمويل النشر - مكافآت", min_value=0.0, value=qvf("تمويل نشر مكافآت مبلغ"), format="%.2f", key="q_prw")
        
        with st.expander("7. الموارد البشرية"):
            c1, c2, c3 = st.columns(3)
            q_per_meet = c1.number_input("اجتماعات دورية", min_value=0, value=qv("عدد الاجتماعات الدورية"), key="q_pm")
            q_board_meet = c2.number_input("مجلس الإدارة", min_value=0, value=qv("عدد اجتماعات مجلس الإدارة"), key="q_bm")
            q_think_meet = c3.number_input("وحدة الفكر", min_value=0, value=qv("عدد اجتماعات وحدة الفكر"), key="q_tm")
            
            q_marketing = st.text_area("آليات التسويق", value=qs("آليات التسويق"), height=80, key="q_mkt")
            q_rewards = st.text_area
